"""Build the editable Quetzal executive PowerPoint with real UI captures."""

from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SCREENSHOTS = ROOT / "docs" / "screenshots" / "deck"
DEFAULT_OUTPUT = ROOT / "docs" / "quetzal-executive-deck.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5

DARK = "071A1B"
DARK_2 = "0B2423"
DARK_3 = "10302E"
GREEN = "0AAE83"
MINT = "73E0C3"
PALE = "EAF6F2"
GOLD = "D8B768"
WHITE = "F7FBFA"
TEXT = "DCEDEA"
MUTED = "8FB0AA"
INK = "10201D"
BLUE = "4A75C5"
RED = "E46E65"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_bg(slide, color: str = DARK) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: str,
    line: str | None = None,
    radius: bool = True,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(1)
    return shape


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = TEXT,
    bold: bool = False,
    font: str = "Arial",
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.03,
    line_spacing: float = 1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_rich_lines(
    slide,
    lines: list[tuple[str, str, bool]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 16,
    gap: float = 6,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    for index, (text, color, bold) in enumerate(lines):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = text
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
    return box


def add_bullets(
    slide,
    items: list[str],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 15,
    color: str = TEXT,
    accent: str = MINT,
    gap: float = 7,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
        r1 = p.add_run()
        r1.text = "• "
        r1.font.name = "Arial"
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = rgb(accent)
        r2 = p.add_run()
        r2.text = item
        r2.font.name = "Arial"
        r2.font.size = Pt(size)
        r2.font.color.rgb = rgb(color)
    return box


def add_kicker(slide, text: str, x: float = 0.62, y: float = 0.25, w: float = 8.0):
    add_text(slide, text.upper(), x, y, w, 0.25, size=8.5, color=MINT, bold=True)


def add_title(
    slide,
    kicker: str,
    title: str,
    subtitle: str | None = None,
    *,
    dark: bool = True,
):
    add_kicker(slide, kicker)
    add_text(
        slide,
        title,
        0.62,
        0.48,
        12.1,
        0.58,
        size=26,
        color=WHITE if dark else INK,
        bold=True,
    )
    if subtitle:
        add_text(
            slide,
            subtitle,
            0.64,
            1.03,
            11.9,
            0.36,
            size=11,
            color=MUTED if dark else "526C66",
        )


def add_footer(slide, number: int, *, dark: bool = True):
    color = MUTED if dark else "647D77"
    add_text(
        slide,
        "QUETZAL · CONEXÃO SEM FRONTEIRAS",
        0.62,
        7.18,
        4.0,
        0.18,
        size=7,
        color=color,
        bold=True,
    )
    add_text(
        slide,
        f"{number:02d}",
        12.28,
        7.16,
        0.42,
        0.2,
        size=8,
        color=color,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )


def add_picture_cover(slide, path: Path, x: float, y: float, w: float, h: float):
    with Image.open(path) as image:
        image_ratio = image.width / image.height
    frame_ratio = w / h
    picture = slide.shapes.add_picture(
        str(path), Inches(x), Inches(y), Inches(w), Inches(h)
    )
    if image_ratio > frame_ratio:
        visible = frame_ratio / image_ratio
        crop = max(0.0, (1.0 - visible) / 2.0)
        picture.crop_left = crop
        picture.crop_right = crop
    else:
        visible = image_ratio / frame_ratio
        crop = max(0.0, (1.0 - visible) / 2.0)
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def add_screenshot(
    slide,
    filename: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    dark: bool = True,
    caption: str | None = None,
):
    border = "2A514B" if dark else "B8CDC7"
    add_rect(slide, x - 0.04, y - 0.04, w + 0.08, h + 0.08, border, border)
    picture = add_picture_cover(slide, SCREENSHOTS / filename, x, y, w, h)
    if caption:
        add_text(
            slide,
            caption,
            x,
            y + h + 0.08,
            w,
            0.25,
            size=8.5,
            color=MUTED if dark else "647D77",
        )
    return picture


def add_metric_card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    label: str,
    value: str,
    detail: str,
    *,
    fill: str = DARK_2,
    value_color: str = WHITE,
):
    add_rect(slide, x, y, w, h, fill, "214641")
    add_text(slide, label.upper(), x + 0.16, y + 0.12, w - 0.32, 0.22, size=8, color=MINT, bold=True)
    add_text(slide, value, x + 0.16, y + 0.39, w - 0.32, 0.38, size=19, color=value_color, bold=True)
    add_text(slide, detail, x + 0.16, y + 0.82, w - 0.32, h - 0.9, size=9.5, color=MUTED)


def new_slide(prs: Presentation, *, color: str = DARK):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, color)
    return slide


def build_deck(output: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "Quetzal — alcance sem fronteiras"
    prs.core_properties.subject = "Quetzal Telecom Command Center"
    prs.core_properties.author = "Quetzal Command Center"
    prs.core_properties.comments = (
        "Deck executivo editável. Capturas do deployment publicado; dados sintéticos."
    )

    # 01 — title
    slide = new_slide(prs)
    add_text(slide, "Q", 0.65, 0.44, 0.55, 0.55, size=25, color=DARK, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    qmark = slide.shapes[-1]
    qmark.fill.solid()
    qmark.fill.fore_color.rgb = rgb(MINT)
    qmark.line.color.rgb = rgb(MINT)
    add_text(slide, "QUETZAL", 1.3, 0.51, 2.3, 0.3, size=15, color=WHITE, bold=True)
    add_text(slide, "CONEXÃO SEM FRONTEIRAS", 1.31, 0.81, 2.6, 0.2, size=7.5, color=MINT, bold=True)
    add_text(slide, "Alcance sem\nfronteiras", 0.7, 1.55, 4.6, 1.55, size=38, color=WHITE, bold=True, line_spacing=0.9)
    add_text(slide, "Da rede ao cliente.\nDo sinal à decisão.", 0.72, 3.25, 4.45, 0.8, size=20, color=MINT, bold=True)
    add_text(slide, "Centro de comando operacional e inteligência governada para telecom em 12 mercados da América Latina.", 0.74, 4.25, 4.25, 1.0, size=13, color=TEXT)
    add_rect(slide, 0.74, 5.65, 4.1, 0.03, GOLD, GOLD, radius=False)
    add_text(slide, "Liberdade sempre · alcance máximo", 0.74, 5.82, 4.4, 0.3, size=11, color="F0D99B", bold=True)
    add_screenshot(slide, "command-center.png", 5.25, 0.48, 7.45, 5.95, caption="Captura real · deployment publicado · dados sintéticos")
    add_footer(slide, 1)

    # 02 — problem and value
    slide = new_slide(prs)
    add_title(slide, "01 · problema telecom", "O sinal se perde entre silos", "Rede, cliente, receita e decisão raramente chegam juntos ao mesmo operador.")
    add_rect(slide, 0.62, 1.55, 6.05, 4.95, DARK_2, "214641")
    add_text(slide, "FRICÇÕES OPERACIONAIS", 0.92, 1.86, 3.2, 0.28, size=10, color=MINT, bold=True)
    add_bullets(slide, [
        "Degradação começa na rede, mas emerge em tickets, NPS, cobrança e churn.",
        "NOC, atendimento, retenção e liderança operam com prioridades diferentes.",
        "Um score isolado não diz quem abordar, com qual oferta ou por quê.",
        "IA sem escopo, evidência e aprovação aumenta o risco operacional.",
    ], 0.9, 2.3, 5.45, 3.35, size=16, gap=12)
    add_rect(slide, 0.92, 5.76, 5.42, 0.47, DARK_3, GOLD)
    add_text(slide, "Da correlação à decisão explicável e executável.", 1.12, 5.88, 5.0, 0.25, size=12, color="F0D99B", bold=True)
    add_metric_card(slide, 6.96, 1.55, 2.68, 1.42, "Prioridade", "Incidente → cliente", "Receita e população expostas no mesmo contexto.")
    add_metric_card(slide, 9.88, 1.55, 2.68, 1.42, "Retenção", "Score → ação", "SLA, NPS, produto, alcance e consentimento.")
    add_metric_card(slide, 6.96, 3.23, 2.68, 1.42, "Campanha", "Audiência → margem", "Capacidade, orçamento e impacto reconciliados.")
    add_metric_card(slide, 9.88, 3.23, 2.68, 1.42, "Inteligência", "Pergunta → prova", "SQL, fontes, memória e guardrails visíveis.")
    add_rect(slide, 6.96, 4.95, 5.60, 1.55, "12332F", "2D5A52")
    add_text(slide, "VALOR", 7.22, 5.2, 1.2, 0.25, size=9, color=MINT, bold=True)
    add_text(slide, "Uma linha contínua entre incidente, ação humana e resultado econômico.", 7.22, 5.55, 5.05, 0.62, size=18, color=WHITE, bold=True)
    add_footer(slide, 2)

    # 03 — journey
    slide = new_slide(prs)
    add_title(slide, "02 · jornada", "Do preflight ao fechamento econômico", "O roteiro demonstra decisão ponta a ponta sem disparar workloads caros ou campanhas.")
    stages = [
        ("01", "Validar", "Warehouse, Genie, IQ, Lakebase, AI/BI e grafo"),
        ("02", "Detectar", "Incidente, cliente e receita potencialmente expostos"),
        ("03", "Explicar", "Churn individual, alcance e exposição separados"),
        ("04", "Decidir", "Voz do Cliente, oferta, consentimento e aprovação"),
        ("05", "Reconciliar", "Custo, benefício, ROI, sensibilidade e limites"),
    ]
    x = 0.63
    for idx, (number, label, detail) in enumerate(stages):
        add_rect(slide, x, 1.55, 2.27, 1.42, DARK_2, "285049")
        add_text(slide, number, x + 0.14, 1.69, 0.38, 0.24, size=10, color=GOLD, bold=True)
        add_text(slide, label, x + 0.14, 1.98, 1.8, 0.3, size=17, color=WHITE, bold=True)
        add_text(slide, detail, x + 0.14, 2.35, 1.95, 0.45, size=9, color=MUTED)
        if idx < len(stages) - 1:
            chevron = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x + 2.31), Inches(1.98), Inches(0.23), Inches(0.5))
            chevron.fill.solid(); chevron.fill.fore_color.rgb = rgb(GREEN); chevron.line.color.rgb = rgb(GREEN)
        x += 2.52
    add_screenshot(slide, "platform-preflight.png", 0.63, 3.4, 3.84, 2.52, caption="Preflight manual")
    add_screenshot(slide, "customer-360.png", 4.76, 3.4, 3.84, 2.52, caption="Cliente 360")
    add_screenshot(slide, "economic-journey.png", 8.89, 3.4, 3.84, 2.52, caption="Fechamento econômico")
    add_text(slide, "Nenhuma campanha é enviada. Geração, Lakeflow, ML e grafo permanecem manuais.", 0.67, 6.3, 12.0, 0.35, size=11, color="F0D99B", bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 3)

    # 04 — architecture
    slide = new_slide(prs)
    add_title(slide, "03 · arquitetura", "Uma plataforma, dois ritmos de dados", "Plano analítico governado no Lakehouse; estado operacional e auditoria no Lakebase.")
    layers = [
        (0.75, 1.7, 2.1, 4.75, "FONTES", ["OSS / BSS", "CRM e billing", "NPS e tickets", "Telemetria e CDR"]),
        (3.05, 1.7, 2.35, 4.75, "LAKEHOUSE", ["Lakeflow", "Unity Catalog", "MLflow", "Gold + triplas"]),
        (5.6, 1.7, 2.35, 4.75, "INTELIGÊNCIA", ["Genie", "AI/BI", "Unity AI Gateway", "Graph + OntoBricks"]),
        (8.15, 1.7, 2.35, 4.75, "APLICAÇÃO", ["FastAPI", "React", "Databricks Apps", "Identidade gerenciada"]),
        (10.7, 1.7, 1.9, 4.75, "ESTADO", ["Lakebase", "Casos", "Campanhas", "Memória + auditoria"]),
    ]
    for idx, (x, y, w, h, label, items) in enumerate(layers):
        fill = DARK_3 if idx in {1, 2, 3} else DARK_2
        add_rect(slide, x, y, w, h, fill, "2B574F")
        add_text(slide, label, x + 0.18, y + 0.18, w - 0.36, 0.3, size=9, color=MINT, bold=True)
        yy = y + 0.78
        for item in items:
            add_rect(slide, x + 0.16, yy, w - 0.32, 0.62, "102B29", "285049")
            add_text(slide, item, x + 0.29, yy + 0.18, w - 0.58, 0.25, size=12, color=WHITE, bold=True)
            yy += 0.82
        if idx < len(layers) - 1:
            chevron = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x + w + 0.05), Inches(3.65), Inches(0.18), Inches(0.55))
            chevron.fill.solid(); chevron.fill.fore_color.rgb = rgb(GREEN); chevron.line.color.rgb = rgb(GREEN)
    add_text(slide, "DAB = código + recursos + promoção reproduzível", 3.35, 6.66, 6.7, 0.3, size=13, color="F0D99B", bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 4)

    # 05 — command center
    slide = new_slide(prs)
    add_title(slide, "04 · experiência operacional", "Centro de mando: decidir antes que o sinal se perca", "A home prioriza sinal, clientes, receita, rede e ações nos 12 mercados.")
    add_screenshot(slide, "command-center.png", 0.63, 1.52, 8.45, 5.15, caption="Centro de mando · tela real do App")
    callouts = [
        ("SINAL PRIORITÁRIO", "Narrativa operacional conecta rede, risco e decisão."),
        ("IMPACTO", "Incidentes críticos e valor mensal em risco na mesma leitura."),
        ("ALCANCE", "KPIs por mercado sem confundir saúde de rede com SLA."),
    ]
    yy = 1.52
    for label, detail in callouts:
        add_rect(slide, 9.45, yy, 3.22, 1.45, DARK_2, "285049")
        add_text(slide, label, 9.7, yy + 0.18, 2.7, 0.24, size=9, color=MINT, bold=True)
        add_text(slide, detail, 9.7, yy + 0.52, 2.68, 0.62, size=13, color=WHITE, bold=True)
        yy += 1.7
    add_footer(slide, 5)

    # 06 — customer 360
    slide = new_slide(prs)
    add_title(slide, "05 · cliente 360", "Cada cliente, sem pontos cegos", "Produto, experiência, SLA, valor, churn e próxima ação em uma visão governada.")
    add_screenshot(slide, "customer-360.png", 0.63, 1.52, 8.45, 5.15, caption="Cliente 360 · identidades de cliente sintéticas")
    add_metric_card(slide, 9.45, 1.52, 3.2, 1.35, "Modelo individual", "Churn ≠ grafo", "A probabilidade ML continua sendo a única probabilidade de churn.")
    add_metric_card(slide, 9.45, 3.05, 3.2, 1.35, "Próxima melhor ação", "Oferta elegível", "Consentimento, produto atual, NPS, SLA e valor orientam a ação.")
    add_metric_card(slide, 9.45, 4.58, 3.2, 1.35, "Ciclo fechado", "Caso + aprovação", "Decisão e execução preservam versões, responsável e auditoria.")
    add_rect(slide, 9.45, 6.1, 3.2, 0.55, "12332F", GOLD)
    add_text(slide, "Sinais separados. Decisão explicável.", 9.62, 6.27, 2.86, 0.22, size=10, color="F0D99B", bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 6)

    # 07 — Genie and IQ
    slide = new_slide(prs)
    add_title(slide, "06 · inteligência governada", "Duas formas de investigar com evidência", "Genie responde sobre dados; Quetzal IQ combina ferramentas curadas, fontes e memória.")
    add_screenshot(slide, "genie-response-chart.png", 0.63, 1.52, 6.0, 3.95, caption="Genie · SQL, resultado e visualização")
    add_screenshot(slide, "quetzal-iq-sources.png", 6.72, 1.52, 6.0, 3.95, caption="Quetzal IQ · pesquisa externa com fontes")
    add_rect(slide, 0.63, 5.83, 6.0, 0.84, DARK_2, "285049")
    add_text(slide, "GENIE", 0.86, 6.03, 0.8, 0.22, size=9, color=MINT, bold=True)
    add_text(slide, "Pergunta → SQL visível → resultado → gráfico interativo", 1.65, 5.98, 4.7, 0.34, size=12.5, color=WHITE, bold=True)
    add_rect(slide, 6.72, 5.83, 6.0, 0.84, DARK_2, "285049")
    add_text(slide, "QUETZAL IQ", 6.95, 6.03, 1.3, 0.22, size=9, color=MINT, bold=True)
    add_text(slide, "Ferramentas internas + Deep Research + fontes + continuidade", 8.2, 5.98, 4.2, 0.34, size=12.5, color=WHITE, bold=True)
    add_footer(slide, 7)

    # 08 — graph
    slide = new_slide(prs)
    add_title(slide, "07 · graph intelligence", "Relacionamentos úteis — sem inventar causalidade", "A rede organiza alcance, exposição e grupos; OntoBricks publica a semântica navegável.")
    add_screenshot(slide, "graph-ontobricks.png", 0.63, 1.52, 8.45, 5.15, caption="Grafo operacional · identidade do operador mascarada")
    lines = [
        ("CHURN INDIVIDUAL", MINT, True),
        ("Probabilidade produzida pelo modelo registrado.", TEXT, False),
        ("EXPOSIÇÃO", MINT, True),
        ("Risco ponderado dos contatos qualificados.", TEXT, False),
        ("ALCANCE", MINT, True),
        ("PageRank ponderado; não mede causalidade.", TEXT, False),
        ("GRUPOS", MINT, True),
        ("Label Propagation para navegação e ação.", TEXT, False),
        ("ONTOBRICKS", MINT, True),
        ("Ontologia, GraphQL e MCP read-only versionados.", TEXT, False),
    ]
    add_rect(slide, 9.4, 1.52, 3.28, 5.15, DARK_2, "285049")
    add_rich_lines(slide, lines, 9.68, 1.78, 2.75, 4.55, size=12.5, gap=4)
    add_footer(slide, 8)

    # 09 — economic journey
    slide = new_slide(prs, color="F1F6F3")
    add_title(slide, "08 · campanhas e economia", "A mesma audiência. As mesmas premissas.", "O App reconcilia população, benefício, custo, impacto, sensibilidade e break-even.", dark=False)
    add_screenshot(slide, "economic-journey.png", 0.63, 1.52, 8.45, 5.15, dark=False, caption="Jornada econômica · identidade do operador mascarada")
    cards = [
        ("ANTES", "Consentimento, serviço antes da venda, capacidade e orçamento."),
        ("DURANTE", "Rascunho versionado, ownership e aprovação segregada."),
        ("DEPOIS", "Resultado observado deve substituir a projeção contrafactual."),
    ]
    yy = 1.52
    for label, detail in cards:
        add_rect(slide, 9.42, yy, 3.25, 1.45, "FFFFFF", "B8CDC7")
        add_text(slide, label, 9.68, yy + 0.18, 2.6, 0.22, size=9, color=GREEN, bold=True)
        add_text(slide, detail, 9.68, yy + 0.52, 2.62, 0.66, size=12.5, color=INK, bold=True)
        yy += 1.7
    add_rect(slide, 9.42, 6.1, 3.25, 0.55, "E4EFEA", GOLD)
    add_text(slide, "Estimativa ≠ resultado realizado", 9.62, 6.26, 2.84, 0.22, size=10, color="735B22", bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 9, dark=False)

    # 10 — governance
    slide = new_slide(prs)
    add_title(slide, "09 · segurança e governança", "A IA não amplia autoridade", "Preflight, identidade, escopo, ferramentas e aprovação permanecem verificáveis.")
    add_screenshot(slide, "platform-preflight.png", 0.63, 1.52, 7.0, 4.58, caption="Preflight manual · 9 dependências prontas")
    add_screenshot(slide, "unity-ai-gateway.png", 7.9, 1.52, 4.8, 2.24, caption="Unity Gateway · modelo governado")
    add_bullets(slide, [
        "Identidade do proxy Databricks Apps; tokens OAuth não persistidos.",
        "RBAC e território avaliados antes de consultar dados.",
        "Quetzal IQ usa ferramentas curadas; não recebe SQL arbitrário.",
        "OntoBricks limitado a intenções MCP read-only.",
        "Solicitante e aprovador são identidades distintas no Lakebase.",
    ], 7.92, 4.18, 4.7, 2.25, size=12.5, gap=7)
    add_footer(slide, 10)

    # 11 — evidence
    slide = new_slide(prs)
    add_title(slide, "10 · evidência real", "A jornada está ancorada em execução", "Históricos reais foram reutilizados; nenhum pipeline, treino ou Genie foi disparado para o deck.")
    add_metric_card(slide, 0.63, 1.55, 2.85, 1.44, "App publicado", "RUNNING", "Compute ACTIVE · deployment SUCCEEDED", value_color=MINT)
    add_metric_card(slide, 3.72, 1.55, 2.85, 1.44, "Lakeflow", "27 / 27", "Run 62176317131697 · SUCCESS", value_color=MINT)
    add_metric_card(slide, 6.81, 1.55, 2.85, 1.44, "Genie", "10 linhas", "Statement 01f1bb5c… · SUCCEEDED", value_color=MINT)
    add_metric_card(slide, 9.90, 1.55, 2.78, 1.44, "Validação", "114 testes", "Ruff · frontend · DAB aprovados", value_color=MINT)
    add_screenshot(slide, "genie-response-chart.png", 0.63, 3.38, 6.1, 3.2, caption="Pergunta real com visualização")
    add_rect(slide, 7.0, 3.38, 5.68, 3.2, DARK_2, "285049")
    add_text(slide, "PERGUNTA COMPROVADA", 7.3, 3.68, 2.7, 0.25, size=9, color=MINT, bold=True)
    add_text(slide, "“Cuales son los clientes con el mayor potencial de abandono?”", 7.3, 4.04, 4.98, 0.78, size=20, color=WHITE, bold=True)
    add_text(slide, "A evidência versionada preserva SQL exato, schema, dez linhas mascaradas e a resposta original do Genie.", 7.3, 5.12, 4.98, 0.72, size=13, color=TEXT)
    add_text(slide, "Lakeflow: update e558c15d… COMPLETED · full refresh manual · sem agenda", 7.3, 6.0, 4.95, 0.36, size=10.5, color="F0D99B", bold=True)
    add_footer(slide, 11)

    # 12 — limitations and contribution
    slide = new_slide(prs)
    add_title(slide, "11 · limites e contribuição", "O que a demo prova — e o que ainda não prova", "Fronteiras explícitas tornam a conversa com o cliente mais confiável.")
    add_rect(slide, 0.63, 1.55, 5.95, 4.95, DARK_2, "285049")
    add_text(slide, "LIMITAÇÕES", 0.93, 1.86, 2.0, 0.26, size=10, color=GOLD, bold=True)
    add_bullets(slide, [
        "Dados, eventos e resultados são sintéticos e demonstrativos.",
        "Scores extremos não provam calibração produtiva.",
        "Relação e grupo indicam associação, nunca causalidade.",
        "Conversão, retenção, receita e ROI são contrafactuais.",
        "Conectores e controles finais dependem dos sistemas do cliente.",
        "Escala, custo, SLAs e qualidade precisam de dados reais.",
    ], 0.92, 2.3, 5.28, 3.75, size=14.5, accent=GOLD, gap=9)
    add_rect(slide, 6.83, 1.55, 5.85, 4.95, DARK_3, "2B574F")
    add_text(slide, "CONTRIBUIÇÃO DO BUILD ATUAL", 7.13, 1.86, 3.4, 0.26, size=10, color=MINT, bold=True)
    add_bullets(slide, [
        "Redesign visual responsivo e identidade Quetzal.",
        "Backend com autenticação, RBAC e falha segura.",
        "Lakeflow/ML recuperados e parametrizados, sem agendas caras.",
        "Cliente 360, SLA, Voz do Cliente e campanhas governadas.",
        "Genie, Quetzal IQ, Lakebase, memória e pesquisa com fontes.",
        "Graph Intelligence + OntoBricks transformados em fluxo útil.",
        "DAB, testes e evidências como base de reconstrução segura.",
    ], 7.12, 2.3, 5.15, 3.9, size=14.5, gap=8)
    add_rect(slide, 3.25, 6.68, 6.8, 0.03, GOLD, GOLD, radius=False)
    add_text(slide, "Um ativo reproduzível, explicável e pronto para receber a realidade do cliente.", 2.2, 6.82, 8.9, 0.28, size=13, color="F0D99B", bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 12)

    # 13 — next step
    slide = new_slide(prs)
    add_title(slide, "12 · próximo passo", "Reconstruir com a realidade do cliente", "A demo vira produto quando dados, políticas e resultados observados substituem as hipóteses.")
    steps = [
        ("01", "Mapear", "OSS, BSS, CRM, billing, NPS, consentimento e contratos."),
        ("02", "Governar", "Identidade, território, papéis, retenção e aprovação."),
        ("03", "Recalibrar", "Labels, features, scores e custos com tempo real observado."),
        ("04", "Homologar", "Ações e ofertas com rede, atendimento, jurídico e segurança."),
        ("05", "Medir", "Uplift controlado e reconciliação do resultado realizado."),
        ("06", "Promover", "DEV → HML → PROD com Databricks Asset Bundles."),
    ]
    yy = 1.55
    for idx, (number, label, detail) in enumerate(steps):
        col_x = 0.68 if idx < 3 else 6.05
        row_y = yy + (idx if idx < 3 else idx - 3) * 1.47
        add_rect(slide, col_x, row_y, 5.0, 1.18, DARK_2, "285049")
        add_text(slide, number, col_x + 0.2, row_y + 0.18, 0.45, 0.26, size=10, color=GOLD, bold=True)
        add_text(slide, label, col_x + 0.75, row_y + 0.14, 1.35, 0.3, size=16, color=WHITE, bold=True)
        add_text(slide, detail, col_x + 0.75, row_y + 0.52, 3.93, 0.42, size=10.5, color=MUTED)
    add_screenshot(slide, "mobile-command-center.png", 11.33, 1.48, 1.42, 4.95, caption="Experiência móvel")
    add_rect(slide, 0.68, 6.2, 10.3, 0.62, "12332F", GOLD)
    add_text(slide, "Quetzal transforma alcance em contexto — e contexto em decisão responsável.", 0.95, 6.37, 9.75, 0.26, size=14, color="F0D99B", bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 13)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build_deck(args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
